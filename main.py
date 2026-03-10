import pandas as pd
import docx
import re

def create_blooket_import(input_docx, output_xlsx):
    # Load the document
    doc = docx.Document(input_docx)
    questions = []
    
    current_question = None
    current_answers = []
    correct_answers = []
    
    # Parse through the word document
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
            
        # Ignore image placeholders (e.g. "[Image 1]") as Blooket imports don't support images this way
        if re.match(r"^\[Image \d+\]$", text, re.IGNORECASE):
            continue
        
        # Check if the text is a question (starts with a number followed by a dot)
        if re.match(r"^\d+\.", text):
            # Save the previous question if it exists
            if current_question is not None:
                questions.append({
                    "question": current_question,
                    "answers": current_answers,
                    "correct": correct_answers
                })
            
            # Remove the "1. " numbering to keep only the question text
            current_question = re.sub(r"^\d+\.\s*", "", text)
            current_answers = []
            correct_answers = []
        else:
            # It's an answer: check if it's the right answer (starts with '*')
            if text.startswith("*"):
                ans_text = text[1:].strip()
                current_answers.append(ans_text)
                correct_answers.append(len(current_answers))  # Save the 1-based index 
            else:
                current_answers.append(text)
                
    # Append the last question
    if current_question is not None:
        questions.append({
            "question": current_question,
            "answers": current_answers,
            "correct": correct_answers
        })
        
    # Prepare the rows for the Blooket format
    rows = []
    for i, q in enumerate(questions, start=1):
        ans1 = q["answers"][0] if len(q["answers"]) > 0 else ""
        ans2 = q["answers"][1] if len(q["answers"]) > 1 else ""
        ans3 = q["answers"][2] if len(q["answers"]) > 2 else ""
        ans4 = q["answers"][3] if len(q["answers"]) > 3 else ""
        
        # Format correct answers as comma separated, e.g., "2" or "1,4"
        correct_str = ",".join(map(str, q["correct"]))
        
        rows.append({
            "Question #": i,
            "Question Text": q["question"],
            "Answer 1": ans1,
            "Answer 2": ans2,
            "Answer 3\n(Optional)": ans3,
            "Answer 4\n(Optional)": ans4,
            "Time Limit (sec)\n(Max: 300 seconds)": 20,  # Default time limit
            "Correct Answer(s)\n(Only include Answer #)": correct_str
        })
        
    df = pd.DataFrame(rows)
    
    # Write to Excel
    with pd.ExcelWriter(output_xlsx, engine='openpyxl') as writer:
        # Start writing headers at row 2 (index 1) to leave room for Blooket's required A1 header
        df.to_excel(writer, index=False, startrow=1)
        
        workbook = writer.book
        worksheet = writer.sheets['Sheet1']
        
        # Add Blooket's required text into the very first cell
        worksheet['A1'] = "Blooket\nImport Template"
        
    print(f"Successfully generated Blooket spreadsheet: {output_xlsx}")

# Run the function
create_blooket_import("quiz.docx", "booklet.xlsx")